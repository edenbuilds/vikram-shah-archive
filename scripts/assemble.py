#!/usr/bin/env python3
"""Assemble transcripts, summaries, Word, and ZIP packs from Firecrawl + pdftotext."""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE_PATH = ROOT / "data" / "archive.json"
ARCHIVE = json.loads(ARCHIVE_PATH.read_text())
FC = ROOT / ".firecrawl"
DL = ROOT / "public" / "downloads"
TR = ROOT / "public" / "transcripts"
GOOGLE_OCR = ROOT / "work" / "google_ocr"
APPLE_OCR = ROOT / "work" / "ocr"
SRC_PDF = Path("/Users/omkar/Downloads/Vikram Shah Documents")

SEAL = RGBColor(0x8B, 0x2E, 0x2E)
INK = RGBColor(0x1C, 0x16, 0x12)
ZEBRA = "F4EEE4"
HEADER_FILL = "8B2E2E"

DISCLAIMER = ARCHIVE["case"]["disclaimer"]
CAUSE = ARCHIVE["case"]["title"]
FORUM = ARCHIVE["case"]["forum"]


def slug_tokens(s: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    return s[:60] or "section"


def pdftotext_pages(pdf: Path, pages: int) -> list[str]:
    out = []
    for i in range(1, pages + 1):
        r = subprocess.run(
            ["pdftotext", "-layout", "-f", str(i), "-l", str(i), str(pdf), "-"],
            capture_output=True,
            text=True,
        )
        out.append((r.stdout or "").replace("\x0c", "").strip())
    return out


def page_ocr_texts(doc_id: str, pages: int) -> list[str] | None:
    """Prefer Google Vision page files, then Apple Vision, when present for every page."""
    for root in (GOOGLE_OCR, APPLE_OCR):
        texts: list[str] = []
        ok = True
        for i in range(1, pages + 1):
            p = root / doc_id / f"page-{i:03d}.txt"
            if not p.exists():
                ok = False
                break
            texts.append(p.read_text(encoding="utf-8", errors="replace").strip())
        if ok and sum(len(t) for t in texts) > 40:
            return texts
    return None


def tag_body(text: str) -> str:
    lines = []
    for raw in text.splitlines():
        line = raw.rstrip()
        low = line.lower()
        if re.search(r"\bsd/?-+\b", line, re.I) or "sd/-" in low:
            lines.append(line)
            lines.append("[SIGNATURE: Wet-ink or sd/- mark as printed on the page.]")
            continue
        if "notary" in low and ("stamp" in low or "seal" in low):
            lines.append(f"[STAMP: {line.strip()}]")
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def split_sections(title: str, body: str, pages: int) -> list[dict]:
    """Split a compilation on annexure / exhibit / divider headings."""
    pattern = re.compile(
        r"(?m)^(#{1,3}\s+.*(?:ANNEXURE|EXHIBIT|APPENDIX|SCHEDULE).*$|^\s*ANNEXURE\s+[-A-Z0-9/]+.*$)",
        re.I,
    )
    matches = list(pattern.finditer(body))
    if len(matches) < 2 or pages < 8:
        return [
            {
                "id": "full",
                "file": "01-full.md",
                "numeral": "I",
                "title": title,
                "short": f"{pages} page{'s' if pages != 1 else ''} as filed.",
                "pages": f"1-{pages}" if pages > 1 else "1",
                "pageStart": 1,
                "pageEnd": pages,
                "mark": "—",
                "body": body,
            }
        ]
    parts = []
    starts = [0] + [m.start() for m in matches]
    # unique ordered starts
    uniq = []
    for s in starts:
        if not uniq or s != uniq[-1]:
            uniq.append(s)
    uniq.append(len(body))
    chunks = []
    for i in range(len(uniq) - 1):
        chunk = body[uniq[i] : uniq[i + 1]].strip()
        if chunk:
            chunks.append(chunk)
    n = len(chunks)
    # page ranges estimated by character share
    total = sum(max(len(c), 1) for c in chunks) or 1
    cursor = 1
    for i, chunk in enumerate(chunks, start=1):
        share = max(1, round(pages * len(chunk) / total))
        start = cursor
        end = pages if i == n else min(pages, cursor + share - 1)
        if end < start:
            end = start
        cursor = end + 1
        head = chunk.splitlines()[0][:120]
        head_clean = re.sub(r"^#+\s*", "", head).strip()
        mark_m = re.search(r"(ANNEXURE|EXHIBIT|APPENDIX)\s*[-:]?\s*([A-Z0-9/]+)", head_clean, re.I)
        mark = mark_m.group(0) if mark_m else ("Pleading" if i == 1 else str(i))
        sid = slug_tokens(head_clean) or f"section-{i:02d}"
        parts.append(
            {
                "id": sid[:48],
                "file": f"{i:02d}-{slug_tokens(head_clean) or 'section'}.md",
                "numeral": str(i),
                "title": head_clean[:140] or f"Section {i}",
                "short": f"Pages {start}-{end} as split from the compilation.",
                "pages": f"{start}-{end}" if start != end else str(start),
                "pageStart": start,
                "pageEnd": end,
                "mark": mark[:40],
                "body": chunk,
            }
        )
    # fix overlapping ids
    seen = set()
    for p in parts:
        base = p["id"]
        k = 2
        while p["id"] in seen:
            p["id"] = f"{base}-{k}"
            k += 1
        seen.add(p["id"])
    return parts


def page_transcript(pages_text: list[str], section_len: int) -> str:
    blocks = []
    for i, body in enumerate(pages_text, start=1):
        tagged = tag_body(body) if body else "[ILLEGIBLE: no extractable text on this page of the PDF; the scan is in Original pages.]"
        blocks.append(f"## Page {i} of {section_len}\n\n{tagged}\n\n---\n")
    return "\n".join(blocks)


def banner(doc: dict) -> str:
    return f"""# {doc['title']}

**Case:** {CAUSE}  
**Forum:** {FORUM}  
**Document:** {doc['title']}  
**Kind:** {doc['kind']}  
**Pages:** {doc['pages']}  
**Source file:** {doc['source']}  
**SHA-256:** `{doc['sha256']}`

> {DISCLAIMER}

"""


def section_file(doc: dict, sec: dict) -> str:
    return f"""# {sec['title']}

**Document:** {doc['title']}  
**Mark:** {sec['mark']}  
**PDF pages:** {sec['pages']}

> {DISCLAIMER}

{sec['body'].strip()}
"""


def full_transcript(doc: dict, sections: list[dict]) -> str:
    rows = ["| No. | Section | PDF pages |", "| --- | --- | --- |"]
    for i, s in enumerate(sections, start=1):
        rows.append(f"| {i} | {s['title']} | {s['pages']} |")
    parts = [banner(doc), "## Contents\n\n" + "\n".join(rows) + "\n"]
    for s in sections:
        parts.append(f"<!-- SECTION: {s['id']} | PDF pages {s['pages']} -->\n")
        parts.append(s["body"].strip() + "\n")
    return "\n".join(parts).strip() + "\n"


def summary_md(doc: dict, sections: list[dict], excerpt: str) -> str:
    map_rows = ["| Pages | Mark | Contents |", "| --- | --- | --- |"]
    for s in sections:
        map_rows.append(f"| {s['pages']} | {s['mark']} | {s['title']} |")
    excerpt = excerpt.strip().split("\n")
    excerpt = "\n".join(excerpt[:40])
    return f"""# Editorial summary — {doc['title']}

**Case:** {CAUSE}  
**Forum:** {FORUM}  
**Document:** {doc['title']}  
**Kind:** {doc['kind']}  
**Pages:** {doc['pages']}

> Editorial only. {DISCLAIMER} Figures below are as stated by the author of the paper, not as found by this archive.

## What this compilation is

1. A paper from the folder **{doc['folder']}**.
2. Classified here as **{doc['kind']}** from the filename and opening pages.
3. Transcribed for reading; the original scan remains the source image.

## Parties / procedural posture

As named on the paper: {CAUSE}. This document is one leaf in that compilation. It is not an award.

## Core money trail or operative facts

Tables and amounts, where they appear, are copied in the transcript. They are not restated as findings here.

## Annexure / document map

{chr(10).join(map_rows)}

## Requested reading / prayers

If the paper contains prayers or a requested order, they remain in the transcript. This summary does not add any.

## Opening extract (as transcribed)

{excerpt}
"""


def set_run_font(run, name="Times New Roman", size=11, italic=False, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_page_number(paragraph):
    run1 = paragraph.add_run("transcription, not a finding  ·  ")
    set_run_font(run1, size=9, italic=True, color=SEAL)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r = paragraph.add_run()
    r._r.append(fld)
    r._r.append(instr)
    r._r.append(fld2)
    set_run_font(r, size=9, color=SEAL)


def write_docx(path: Path, doc: dict, markdown: str) -> None:
    d = Document()
    for section in d.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        hp = section.header.paragraphs[0]
        hp.text = CAUSE
        for run in hp.runs:
            set_run_font(run, size=9, italic=True, color=SEAL)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_number(fp)

    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.font.color.rgb = INK
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(ARCHIVE["case"]["kicker"].upper())
    set_run_font(r, size=12, color=SEAL, bold=True)
    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(doc["title"])
    set_run_font(r, size=18, bold=True)
    sub = d.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(CAUSE)
    set_run_font(r, size=12, italic=True)
    disc = d.add_paragraph()
    disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = disc.add_run(DISCLAIMER)
    set_run_font(r, size=11, italic=True, color=SEAL)
    d.add_page_break()

    # strip illegal XML chars
    markdown = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", markdown)

    def add_tagged(text: str):
        p = d.add_paragraph()
        r = p.add_run(text)
        set_run_font(r, size=11, italic=True, color=SEAL)

    i = 0
    lines = markdown.splitlines()
    while i < len(lines):
        line = lines[i]
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*-+", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|?\s*-+", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = d.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(cols):
                        cell = table.rows[ri].cells[ci]
                        cell.text = row[ci] if ci < len(row) else ""
                        tcPr = cell._tc.get_or_add_tcPr()
                        shading = OxmlElement("w:shd")
                        fill = HEADER_FILL if ri == 0 else (ZEBRA if ri % 2 == 0 else "FFFFFF")
                        shading.set(qn("w:fill"), fill)
                        shading.set(qn("w:val"), "clear")
                        tcPr.append(shading)
                        for p in cell.paragraphs:
                            for run in p.runs:
                                set_run_font(run, size=10, bold=(ri == 0), color=(RGBColor(255, 255, 255) if ri == 0 else INK))
            continue
        m = re.match(r"^\[([A-Z/ ]+):\s*(.*)\]\s*$", line)
        if m:
            add_tagged(f"{m.group(1)} {m.group(2)}")
            i += 1
            continue
        if line.startswith("# "):
            p = d.add_paragraph()
            r = p.add_run(line[2:].strip())
            set_run_font(r, size=16, bold=True, color=SEAL)
        elif line.startswith("## "):
            p = d.add_paragraph()
            r = p.add_run(line[3:].strip())
            set_run_font(r, size=13, bold=True, color=SEAL)
        elif line.startswith("### "):
            p = d.add_paragraph()
            r = p.add_run(line[4:].strip())
            set_run_font(r, size=12, bold=True)
        elif line.startswith("> "):
            add_tagged(line[2:])
        elif line.strip() == "---":
            d.add_paragraph()
        elif line.strip():
            p = d.add_paragraph()
            # bold/italic leftovers
            r = p.add_run(re.sub(r"[*_]", "", line))
            set_run_font(r, size=11)
        i += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)


def case_summary(docs_done: list[dict]) -> str:
    rows = ["| Bundle | Document | Kind | Pages |", "| --- | --- | --- | --- |"]
    for d in ARCHIVE["docs"]:
        rows.append(f"| {d['bundleTitle']} | {d['title']} | {d['kind']} | {d['pages']} |")
    return f"""# Case summary — {CAUSE}

**Forum:** {FORUM}  
**Venue:** {ARCHIVE['case']['venue']}  
**Papers in this archive:** {ARCHIVE['case']['docCount']} PDF files, {ARCHIVE['case']['pageCount']} pages.

> Editorial only. {DISCLAIMER}

## What this compilation is

1. A public reading copy of papers from the Vikram Shah document tree, kept as a case rather than a single PDF.
2. Each source file is its own document, with its own transcript, original scans, and downloads.
3. Related High Court appointment papers, arbitral pleadings, minutes, financial statements, and civil-suit extracts sit beside one another so a stranger can see how the authors connected them.

## Parties / procedural posture

- Claimant, as named: {ARCHIVE['case']['claimant']}
- Respondents 1–3, as named: {ARCHIVE['case']['respondents13']}
- Respondent No. 4, as named: {ARCHIVE['case']['respondent4']}
- Appointment proceeding: AAR No. 13 of 2024, High Court of Bombay at Goa, as stated on those papers.
- Related civil proceeding: SCS 52/2013/A, as stated on those papers.

## Core money trail or operative facts

Amounts, dates, and account numbers appear in the financial documents and pleadings. They are copied in those transcripts. This page does not choose among competing figures.

## Document map

{chr(10).join(rows)}

## Requested reading / prayers

Prayers remain in the pleadings that contain them. This summary does not add a requested outcome.
"""


def assemble_doc(doc: dict) -> dict | None:
    fc = FC / f"{doc['id']}.md"
    pdf = DL / doc["file"]
    pages_text = pdftotext_pages(pdf, doc["pages"])
    text_chars = sum(len(p) for p in pages_text)
    text_ok = text_chars > 200 * max(doc["pages"], 1) * 0.3
    fc_body = ""
    if fc.exists() and fc.stat().st_size > 80:
        fc_body = fc.read_text(encoding="utf-8", errors="replace")
    vision_pages = page_ocr_texts(doc["id"], doc["pages"])
    vision_chars = sum(len(p) for p in vision_pages) if vision_pages else 0
    # Prefer the denser faithful source. Page-keyed extracts win for legal navigation.
    if vision_pages and vision_chars >= max(text_chars, len(fc_body) if fc_body else 0) * 0.85:
        body = page_transcript(vision_pages, doc["pages"])
        source = "google-vision" if (GOOGLE_OCR / doc["id"]).exists() else "apple-vision"
    elif text_ok and (not fc_body or text_chars >= len(fc_body) * 0.9):
        body = page_transcript(pages_text, doc["pages"])
        source = "pdftotext"
    elif fc_body:
        body = fc_body
        source = "firecrawl"
        # If Firecrawl returned a thin continuous blob, prefer page-keyed Vision/pdftotext.
        if vision_pages and (len(fc_body) < doc["pages"] * 900 or not re.search(r"(?mi)^## Page\s+\d+", fc_body)):
            if vision_chars >= len(fc_body) * 0.7 or len(fc_body) < doc["pages"] * 700:
                body = page_transcript(vision_pages, doc["pages"])
                source = "google-vision" if (GOOGLE_OCR / doc["id"]).exists() else "apple-vision"
            else:
                body = (
                    fc_body.strip()
                    + "\n\n---\n\n## Page-keyed extract (Vision OCR)\n\n"
                    + page_transcript(vision_pages, doc["pages"])
                )
                source = "firecrawl+vision"
        elif text_chars > 500 and len(fc_body) < doc["pages"] * 350:
            body = (
                fc_body.strip()
                + "\n\n---\n\n## Page-keyed extract (pdftotext)\n\n"
                + page_transcript(pages_text, doc["pages"])
            )
            source = "firecrawl+pdftotext"
    elif text_ok:
        body = page_transcript(pages_text, doc["pages"])
        source = "pdftotext"
    elif vision_pages:
        body = page_transcript(vision_pages, doc["pages"])
        source = "google-vision" if (GOOGLE_OCR / doc["id"]).exists() else "apple-vision"
    else:
        return None
    body = tag_body(body)
    if source.startswith("firecrawl") and not re.search(r"(?mi)^## Page\s+\d+", body):
        body = (
            f"> Clerk note: continuous OCR for a {doc['pages']}-page paper. "
            f"Open Original scans for page boundaries. Figures are as printed.\n\n"
            + body
        )
    sections = split_sections(doc["title"], body, doc["pages"])
    # Page-keyed sources keep the page transcript as the single section body when unsplit.
    if source in ("pdftotext", "google-vision", "apple-vision") or source.endswith("+vision"):
        if len(sections) == 1:
            sections[0]["body"] = body

    tdir = TR / doc["id"] / "sections"
    ddir = DL / doc["id"] / "sections"
    tdir.mkdir(parents=True, exist_ok=True)
    ddir.mkdir(parents=True, exist_ok=True)

    full = full_transcript(doc, sections)
    summary = summary_md(doc, sections, body[:2500])
    (TR / doc["id"] / "FULL-TRANSCRIPT.md").write_text(full, encoding="utf-8")
    (DL / doc["id"] / "FULL-TRANSCRIPT.md").write_text(full, encoding="utf-8")
    (TR / doc["id"] / "SUMMARY.md").write_text(summary, encoding="utf-8")
    (DL / doc["id"] / "SUMMARY.md").write_text(summary, encoding="utf-8")
    shutil.copy2(TR / doc["id"] / "FULL-TRANSCRIPT.md", DL / f"{doc['id']}-FULL-TRANSCRIPT.md")
    shutil.copy2(TR / doc["id"] / "SUMMARY.md", DL / f"{doc['id']}-SUMMARY.md")

    for s in sections:
        text = section_file(doc, s)
        (tdir / s["file"]).write_text(text, encoding="utf-8")
        (ddir / s["file"]).write_text(text, encoding="utf-8")
        shutil.copy2(tdir / s["file"], DL / f"{doc['id']}-{s['file']}")

    docx_path = DL / doc["id"] / "FULL-TRANSCRIPT.docx"
    write_docx(docx_path, doc, full)
    shutil.copy2(docx_path, DL / f"{doc['id']}-FULL-TRANSCRIPT.docx")

    zip_path = DL / f"{doc['id']}-transcripts.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(DL / doc["id"] / "FULL-TRANSCRIPT.md", "FULL-TRANSCRIPT.md")
        z.write(docx_path, "FULL-TRANSCRIPT.docx")
        z.write(DL / doc["id"] / "SUMMARY.md", "SUMMARY.md")
        for s in sections:
            z.write(ddir / s["file"], f"sections/{s['file']}")
        z.write(pdf, doc["file"])

    clean_sections = [{k: v for k, v in s.items() if k != "body"} for s in sections]
    doc["sections"] = clean_sections
    doc["transcribed"] = True
    doc["transcriptSource"] = source
    doc["note"] = sections[0]["short"] if len(sections) == 1 else f"{len(sections)} sections as split from the paper."
    return doc


def main() -> None:
    updated = []
    missing = []
    for doc in ARCHIVE["docs"]:
        res = assemble_doc(doc)
        if res:
            updated.append(res)
            print(f"ok {doc['id']} ({doc['pages']}p, {len(doc['sections'])} sections)", flush=True)
        else:
            missing.append(doc["id"])
            doc["transcribed"] = False
            print(f"pending {doc['id']}", flush=True)

    ARCHIVE["docs"] = [d for d in ARCHIVE["docs"]]
    figures = [
        {"label": "Papers", "value": str(ARCHIVE["case"]["docCount"]), "note": "Separate PDFs, each a document"},
        {"label": "Pages", "value": str(ARCHIVE["case"]["pageCount"]), "note": "Original scans in this archive"},
        {"label": "Forum", "value": "AAR 13/2024", "note": "As stated on the High Court appointment papers"},
        {"label": "Civil suit", "value": "SCS 52/2013/A", "note": "Communidade of Bambolim papers as filed"},
        {"label": "Loonkar confirmation", "value": "Rs. 3.676 crore", "note": "Letter dated 28.12.11, as printed"},
        {"label": "Transcribed so far", "value": str(len(updated)), "note": "Documents with a full typed file"},
    ]
    ARCHIVE["figures"] = figures
    ARCHIVE_PATH.write_text(json.dumps(ARCHIVE, indent=2), encoding="utf-8")
    (ROOT / "public" / "archive.json").write_text(json.dumps(ARCHIVE, indent=2), encoding="utf-8")

    cs = case_summary(updated)
    (TR / "CASE-SUMMARY.md").write_text(cs, encoding="utf-8")
    (DL / "CASE-SUMMARY.md").write_text(cs, encoding="utf-8")

    # master zip of transcripts only (no page JPEGs / no original PDFs — those stay on Supabase)
    master = DL / "shah-v-trindade-archive.zip"
    with zipfile.ZipFile(master, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(DL / "CASE-SUMMARY.md", "CASE-SUMMARY.md")
        for doc in ARCHIVE["docs"]:
            if not doc.get("transcribed"):
                continue
            prefix = doc["id"]
            z.write(DL / f"{prefix}-FULL-TRANSCRIPT.md", f"{prefix}/FULL-TRANSCRIPT.md")
            z.write(DL / f"{prefix}-SUMMARY.md", f"{prefix}/SUMMARY.md")
            docx = DL / f"{prefix}-FULL-TRANSCRIPT.docx"
            if docx.exists():
                z.write(docx, f"{prefix}/FULL-TRANSCRIPT.docx")
            for s in doc["sections"]:
                p = DL / f"{prefix}-{s['file']}"
                if p.exists():
                    z.write(p, f"{prefix}/sections/{s['file']}")
    print("missing", len(missing), "done", len(updated), "master", master.stat().st_size)
    if missing:
        print("pending ids sample:", ", ".join(missing[:12]))


if __name__ == "__main__":
    main()
